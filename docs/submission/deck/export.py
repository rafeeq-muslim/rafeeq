import pymupdf, io, glob, os
from pptx import Presentation
from pptx.util import Emu
for f in glob.glob('/tmp/claude-1000/s3-sub/pages/s*.png'): os.remove(f)
d=pymupdf.open('rafeeq-deck.vector.pdf')
prs=Presentation(); prs.slide_width=Emu(12192000); prs.slide_height=Emu(6858000)
out=pymupdf.open()
for i,p in enumerate(d):
    z=1920/p.rect.width
    pix=p.get_pixmap(matrix=pymupdf.Matrix(z,z), alpha=False)
    pix.save(f'/tmp/claude-1000/s3-sub/pages/s{i+1:02d}.png')
    jpg=pix.tobytes('jpeg', jpg_quality=90)
    s=prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(io.BytesIO(jpg),0,0,width=prs.slide_width,height=prs.slide_height)
    s.notes_slide.notes_text_frame.text=p.get_text().strip()
    pg=out.new_page(width=1440,height=810); pg.insert_image(pg.rect, stream=jpg)
prs.save('rafeeq-deck.pptx'); out.save('rafeeq-deck.pdf', deflate=True)
print(len(d))
